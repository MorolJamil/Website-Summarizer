"""
Website Page Summarizer
-----------------------
Summarizes ONE web page (only the page at the link you give, not the whole site)
and saves the result as an Excel table in a separate "summaries" folder.

HOW TO USE:
  1. Change PAGE_URL below to any page you want.
  2. (Optional) Change SAVE_FOLDER to choose where the file is saved.
  3. Run the script (in VS Code: click the Run button, or press Ctrl+F5).
  4. The Excel file is saved in Downloads/Website Summaries (or your folder)
     and opens automatically.

First-time setup (run once in the VS Code terminal):
  pip install requests beautifulsoup4 openpyxl
"""

# ======================================================================
#  1) CHANGE THE PAGE LINK HERE
PAGE_URL = "https://www.bbc.com/news"

#  2) CHOOSE WHERE THE SUMMARY IS SAVED (optional)
#     Leave it empty ("") to save into your Downloads folder, inside
#     a folder called "Website Summaries".
#     Or type any folder path, for example:
#       Windows: r"C:\Users\YourName\Documents\My Summaries"
#       Mac:     "/Users/YourName/Documents/My Summaries"
SAVE_FOLDER = ""
# ======================================================================

import re
from collections import Counter
from datetime import datetime
from pathlib import Path
from urllib.parse import urlparse

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

# Output folder: your chosen folder, or Downloads/Website Summaries by default
if SAVE_FOLDER.strip():
    OUTPUT_FOLDER = Path(SAVE_FOLDER.strip()).expanduser()
else:
    OUTPUT_FOLDER = Path.home() / "Downloads" / "Website Summaries"

MAX_SECTIONS = 10          # max number of page sections listed in the table
OVERVIEW_SENTENCES = 5     # sentences in the overall page description
SECTION_SENTENCES = 2      # sentences in each section description

STOPWORDS = set("""
a about above after again against all am an and any are as at be because been
before being below between both but by can could did do does doing down during
each few for from further had has have having he her here hers herself him
himself his how i if in into is it its itself just let me more most my myself
no nor not now of off on once only or other our ours ourselves out over own
same she should so some such than that the their theirs them themselves then
there these they this those through to too under until up very was we were
what when where which while who whom why will with would you your yours
yourself yourselves also get us new one use using may like
""".split())


# ---------------------------------------------------------------- fetching
def fetch_page(url: str) -> BeautifulSoup:
    headers = {
        "User-Agent": (
            "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
            "(KHTML, like Gecko) Chrome/124.0 Safari/537.36"
        ),
        "Accept-Language": "en-US,en;q=0.9",
    }
    response = requests.get(url, headers=headers, timeout=20)
    response.raise_for_status()
    response.encoding = response.apparent_encoding or response.encoding
    return BeautifulSoup(response.text, "html.parser")


def clean(text: str) -> str:
    return re.sub(r"\s+", " ", text or "").strip()


# ---------------------------------------------------------------- extracting
def get_main_content(soup: BeautifulSoup):
    """Remove invisible and repeated parts (scripts, menus, footers)."""
    for tag in soup(["script", "style", "noscript", "svg", "iframe",
                     "template", "form", "button"]):
        tag.decompose()
    # Remove elements hidden from view
    for tag in soup.select('[hidden], [aria-hidden="true"], [style*="display:none"], '
                           '[style*="display: none"]'):
        tag.decompose()

    content = soup.find("main") or soup.find("article") or soup.body or soup
    for tag in content.find_all(["nav", "footer", "aside"]):
        tag.decompose()
    return content


def split_sections(content):
    """Group the page text under its headings (h1/h2/h3)."""
    sections = []
    current = {"heading": "Introduction", "text": []}
    seen = set()

    for el in content.find_all(["h1", "h2", "h3", "p", "li", "blockquote", "td"]):
        text = clean(el.get_text(" "))
        if not text or text in seen:
            continue
        seen.add(text)
        if el.name in ("h1", "h2", "h3"):
            if current["text"] or current["heading"] != "Introduction":
                sections.append(current)
            current = {"heading": text[:120], "text": []}
        else:
            current["text"].append(text)
    sections.append(current)

    # Fallback for pages built without <p>/<li> tags
    if not any(s["text"] for s in sections):
        lines = [clean(l) for l in content.get_text("\n").split("\n")]
        sections = [{"heading": "Page content",
                     "text": [l for l in lines if len(l) > 20]}]
    return sections


# ---------------------------------------------------------------- summarizing
def to_sentences(texts):
    joined = " ".join(t if t[-1] in ".!?" else t + "." for t in texts if t)
    parts = re.split(r"(?<=[.!?])\s+(?=[A-Z0-9\"'])", joined)
    sentences = [p.strip() for p in parts if 30 <= len(p.strip()) <= 400]
    if not sentences:  # short snippets without punctuation (common on home pages)
        sentences = [t for t in texts if len(t) >= 15]
    return sentences


def word_list(text):
    return [w for w in re.findall(r"[a-zA-Z][a-zA-Z'-]+", text.lower())
            if w not in STOPWORDS and len(w) > 2]


def summarize(texts, n_sentences, frequencies=None):
    """Pick the most important sentences, keeping their original order."""
    sentences = to_sentences(texts)
    if not sentences:
        return ""
    if frequencies is None:
        frequencies = Counter(word_list(" ".join(texts)))
    if not frequencies:
        return " ".join(sentences[:n_sentences])
    top = max(frequencies.values())

    scored = []
    for i, s in enumerate(sentences):
        words = word_list(s)
        if not words:
            continue
        score = sum(frequencies[w] / top for w in words) / (len(words) ** 0.5)
        scored.append((score, i, s))

    best = sorted(scored, reverse=True)[:n_sentences]
    best.sort(key=lambda x: x[1])
    return " ".join(s if s[-1] in ".!?" else s + "." for _, _, s in best)


def shorten(text, max_words=25):
    words = text.split()
    return text if len(words) <= max_words else " ".join(words[:max_words]) + "..."


# ---------------------------------------------------------------- building rows
def build_rows(url):
    soup = fetch_page(url)

    title = clean(soup.title.get_text()) if soup.title else ""
    meta = soup.find("meta", attrs={"name": re.compile("^description$", re.I)}) \
        or soup.find("meta", attrs={"property": "og:description"})
    meta_description = clean(meta.get("content", "")) if meta else ""

    content = get_main_content(soup)
    sections = split_sections(content)
    all_text = [t for s in sections for t in s["text"]]
    page_frequencies = Counter(word_list(" ".join(all_text)))

    rows = []

    # Row 1: overall page summary
    overview = summarize(all_text, OVERVIEW_SENTENCES, page_frequencies)
    short = meta_description or summarize(all_text, 1, page_frequencies) or title
    short_label = f"{title}: {shorten(short)}" if title else shorten(short)
    description = overview or meta_description or "No readable text found on this page."
    rows.append((short_label, description))

    # Following rows: one per section of the page
    for section in sections:
        if len(rows) > MAX_SECTIONS:
            break
        if not section["text"]:
            continue
        desc = summarize(section["text"], SECTION_SENTENCES, page_frequencies)
        if not desc:
            desc = shorten(" ".join(section["text"]), 60)
        rows.append((section["heading"], desc))

    return title, rows


# ---------------------------------------------------------------- saving
def save_table(url, title, rows):
    OUTPUT_FOLDER.mkdir(parents=True, exist_ok=True)
    domain = urlparse(url).netloc.replace("www.", "").replace(":", "_") or "page"
    stamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
    file_path = OUTPUT_FOLDER / f"{domain}_summary_{stamp}.xlsx"

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"

    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    base_font = Font(name="Arial", size=11)

    ws.append(["Short Summary", "Description"])
    for cell in ws[1]:
        cell.font = Font(name="Arial", size=12, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    for short, desc in rows:
        ws.append([short, desc])
        for cell in ws[ws.max_row]:
            cell.font = base_font
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = border
    for cell in ws[2]:  # highlight the overall summary row
        cell.font = Font(name="Arial", size=11, bold=True)
        cell.fill = PatternFill("solid", fgColor="DDEBF7")

    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 100
    ws.freeze_panes = "A2"

    # Source info under the table
    info_row = ws.max_row + 2
    ws.cell(info_row, 1, "Source page:").font = Font(name="Arial", size=10, italic=True)
    ws.cell(info_row, 2, url).font = Font(name="Arial", size=10, italic=True)
    ws.cell(info_row + 1, 1, "Created:").font = Font(name="Arial", size=10, italic=True)
    ws.cell(info_row + 1, 2, datetime.now().strftime("%Y-%m-%d %H:%M")).font = \
        Font(name="Arial", size=10, italic=True)

    wb.save(file_path)
    return file_path


# ---------------------------------------------------------------- main
def main():
    url = PAGE_URL.strip()
    if not url.startswith(("http://", "https://")):
        url = "https://" + url

    print(f"Reading: {url}")
    try:
        title, rows = build_rows(url)
    except requests.RequestException as error:
        print(f"Could not open the page: {error}")
        return

    print(f"\nPage title: {title or '(none)'}\n")
    for short, desc in rows:
        print(f"- {short}\n    {desc}\n")

    try:
        file_path = save_table(url, title, rows)
    except OSError as error:
        print(f"Could not save to '{OUTPUT_FOLDER}': {error}")
        print("Check the SAVE_FOLDER path at the top of the script.")
        return
    print(f"Table saved to: {file_path}")
    open_file(file_path)


def open_file(file_path):
    """Open the saved Excel file with the computer's default program."""
    import os
    import subprocess
    import sys
    try:
        if sys.platform.startswith("win"):
            os.startfile(file_path)
        elif sys.platform == "darwin":
            subprocess.run(["open", str(file_path)], check=False)
        else:
            subprocess.run(["xdg-open", str(file_path)], check=False)
    except Exception:
        pass  # If it can't open automatically, the file is still saved


if __name__ == "__main__":
    main()